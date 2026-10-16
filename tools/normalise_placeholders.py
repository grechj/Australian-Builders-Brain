"""Rename legacy placeholders in the Word templates to the canonical field names.

Usage:
    python tools/normalise_placeholders.py            # report only
    python tools/normalise_placeholders.py --write    # rewrite templates in place

Canonical names are listed in brain/fields/project_fields.md. Add a row to
RENAMES whenever an imported template uses a different spelling.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).parent))
from docx_tools import replace_in_document  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent

RENAMES = {
    "[YOUR LOGO HERE]": "[LOGO]",
    "[BUSINESS NAME]": "[BUSINESS_NAME]",
    "[YOUR ABN]": "[ABN]",
    "[LICENCE NUMBER]": "[LICENCE_NUMBER]",
    "[ADDRESS]": "[BUSINESS_ADDRESS]",
    "[ADDRESS LINE 1]": "[BUSINESS_ADDRESS]",
    "[ADDRESS LINE 2]": "[BUSINESS_ADDRESS_LINE_2]",
    "[YOUR WEBSITE]": "[WEBSITE]",
    "[CLIENT NAME]": "[CLIENT_NAME]",
    "[CLIENT ADDRESS]": "[CLIENT_ADDRESS]",
    "[CONTRACT NUMBER]": "[CONTRACT_NUMBER]",
    "[ORIGINAL CONTRACT SUM]": "[CONTRACT_SUM]",
    "[AUTO OR MANUAL]": "[CONTRACT_CLAUSE]",
    "[PROJECT NAME/ADDRESS]": "[PROJECT_ADDRESS]",
    "[PROJECT NAME]": "[PROJECT_NAME]",
    "[QUOTE NUMBER]": "[QUOTE_NUMBER]",
    "[TRADE/NAME]": "[TRADE_NAME]",
    "[HOSPITAL_NAME & ADDRESS]": "[NEAREST_HOSPITAL]",
    "[VAR-001]": "[VAR_NUMBER]",
    "[CR-001]": "[CHANGE_REQUEST_NUMBER]",
    "[BC-001]": "[BACK_CHARGE_NUMBER]",
    "[DELAY-001]": "[DELAY_NUMBER]",
    "[INV-001]": "[INVOICE_NUMBER]",
    "[SWMS-001]": "[SWMS_NUMBER]",
    "[SI-001]": "[SITE_INSTRUCTION_NUMBER]",
}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--write", action="store_true", help="save changes")
    args = parser.parse_args()

    total = 0
    for path in sorted((ROOT / "templates" / "word").glob("*.docx")):
        doc = Document(path)
        n = replace_in_document(doc, RENAMES.get)
        if n:
            total += n
            print(f"{path.name}: {n} renamed")
            if args.write:
                doc.save(path)
    print(f"{total} placeholders {'renamed' if args.write else 'would be renamed'}")


if __name__ == "__main__":
    main()
