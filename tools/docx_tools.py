"""Shared helpers for reading and rewriting placeholders in .docx templates.

Word often splits one visible string such as "[CLIENT_NAME]" across several
runs (for example when part of it was retyped or spell-checked). Replacing
run by run would miss those, so we replace across the whole paragraph and
write the result back into the runs, keeping the formatting of the run where
each placeholder starts.
"""

from __future__ import annotations

import re
from typing import Callable, Iterator

from docx.document import Document as DocumentType
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

# A placeholder is a bracketed upper-case token: [CLIENT_NAME], [VAR-001],
# [DD/MM/YYYY]. Bracketed prose like "[Describe the variation...]" is not.
PLACEHOLDER = re.compile(r"\[[A-Z0-9][A-Z0-9 _/&.-]*\]")


def _table_paragraphs(table: Table) -> Iterator[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for inner in cell.tables:
                yield from _table_paragraphs(inner)


def iter_paragraphs(doc: DocumentType) -> Iterator[Paragraph]:
    """Every paragraph in the body, tables, headers and footers."""
    yield from doc.paragraphs
    for table in doc.tables:
        yield from _table_paragraphs(table)
    for section in doc.sections:
        for part in (section.header, section.footer,
                     section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            if part.is_linked_to_previous:
                continue
            yield from part.paragraphs
            for table in part.tables:
                yield from _table_paragraphs(table)


def find_placeholders(doc: DocumentType) -> set[str]:
    found = set()
    for p in iter_paragraphs(doc):
        found.update(PLACEHOLDER.findall("".join(r.text for r in p.runs)))
    return found


def replace_in_paragraph(p: Paragraph, replace: Callable[[str], str | None],
                         pattern: re.Pattern = PLACEHOLDER) -> int:
    """Replace placeholders (or any `pattern` match) in one paragraph.

    `replace` receives a placeholder such as "[CLIENT_NAME]" and returns the
    new text, or None to leave it alone. Returns the number of replacements.
    """
    runs = p.runs
    texts = [r.text for r in runs]
    starts, pos = [], 0
    for t in texts:
        starts.append(pos)
        pos += len(t)
    full = "".join(texts)

    def run_at(offset: int) -> int:
        # Text left of the match being replaced is never modified, so the
        # original start offsets still locate it.
        for i, start in enumerate(starts):
            if start <= offset < start + len(texts[i]):
                return i
        raise ValueError(f"offset {offset} outside paragraph")

    count = 0
    # Right to left, so earlier offsets stay valid as later text changes.
    for m in reversed(list(pattern.finditer(full))):
        new = replace(m.group(0))
        if new is None:
            continue
        s, e = m.span()
        i, j = run_at(s), run_at(e - 1)
        head = texts[i][: s - starts[i]]
        if i == j:
            texts[i] = head + new + texts[i][e - starts[i]:]
        else:
            texts[i] = head + new
            for k in range(i + 1, j):
                texts[k] = ""
            texts[j] = texts[j][e - starts[j]:]
        count += 1

    if count:
        for r, t in zip(runs, texts):
            if r.text != t:
                r.text = t
    return count


def replace_in_document(doc: DocumentType, replace: Callable[[str], str | None],
                        pattern: re.Pattern = PLACEHOLDER) -> int:
    return sum(replace_in_paragraph(p, replace, pattern) for p in iter_paragraphs(doc))


def iter_cells(doc: DocumentType) -> Iterator[list[_Cell]]:
    """Each table row as a list of its distinct cells, in document order.

    Merged cells appear once per row, not once per grid column.
    """
    def walk(table: Table):
        for row in table.rows:
            cells, seen = [], set()
            for cell in row.cells:
                if id(cell._tc) not in seen:
                    seen.add(id(cell._tc))
                    cells.append(cell)
            yield cells
            for cell in cells:
                for inner in cell.tables:
                    yield from walk(inner)
    for table in doc.tables:
        yield from walk(table)


def cell_text(cell: _Cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs)


def set_cell_text(cell: _Cell, text: str) -> None:
    """Replace a cell's text, keeping the formatting of its first run.

    Newlines in `text` become separate paragraphs.
    """
    first = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    lines = text.split("\n")
    runs = first.runs
    if runs:
        runs[0].text = lines[0]
        for r in runs[1:]:
            r.text = ""
    else:
        first.add_run(lines[0])
    for line in lines[1:]:
        para = cell.add_paragraph(style=first.style)
        para.paragraph_format.alignment = first.paragraph_format.alignment
        run = para.add_run(line)
        if runs:
            run.bold, run.italic = runs[0].bold, runs[0].italic
            run.font.size, run.font.name = runs[0].font.size, runs[0].font.name
