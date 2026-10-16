"""Check the brain/ bundle against OKF v0.2 and this repo's own rules.

Usage:
    python tools/okf_lint.py                   # check, exit 1 on errors
    python tools/okf_lint.py --write_indexes   # regenerate every index.md, then check

Errors: missing frontmatter or `type`, frontmatter in a non-root index.md,
footnotes with no matching source, template placeholders not in the field
schema, computation hash mismatches, spaces or hyphens in file names.
Warnings: broken links (OKF allows them, but they mark missing knowledge)
and concepts past `stale_after`.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parent.parent
BRAIN = ROOT / "brain"
RESERVED = {"index.md", "log.md"}
OKF_VERSION = "0.2"

FOLDERS = {
    "lifecycle": "The five stages of a job, from quote to final account, with the documents used at each",
    "phases": "Construction phases by project type: activities, inspections, materials, weather risk, next trades",
    "contracts": "Contract topics: variations, extensions of time, HIA/MBA clause map",
    "states": "State and territory rules: regulators, home building and security-of-payment laws",
    "ncc": "National Construction Code: editions, volumes, how to cite",
    "documents": "One page per template: what it's for, its fields, which skill fills it",
    "fields": "The shared project fields every template fills from",
    "computations": "Attested money calculations: GST, variation price, margin",
    "glossary": "Site slang and trade terms",
    "references": "Playbooks, sanctioned code and attesters that other concepts point to",
    "new_home": "New home construction phases",
    "renovation": "Renovation construction phases",
    "extension": "Extension construction phases",
    "commercial": "Commercial construction phases",
    "attesters": "Deterministic checks for computation receipts",
    "2025": "NCC 2025 (v1.2), converted from the ABCB's CC BY 4.0 XML: every Part, clause, definition and state variation",
    "volume_one": "Class 2 to 9 buildings: apartments, commercial, industrial and public buildings",
    "volume_two": "Class 1 and 10 buildings: houses, townhouses, garages, sheds, pools",
    "volume_three": "Plumbing and drainage",
    "housing_provisions": "ABCB Housing Provisions: deemed-to-satisfy construction details for houses",
    "livable_housing_design": "ABCB Livable Housing Design Standard",
}
TITLES = {"ncc": "NCC", "2025": "NCC 2025"}
ROOT_INTRO = """The Australian Builder's Brain: knowledge for Australian builders and
tradespeople, from quote to final account. Start with the job's
[lifecycle stage](lifecycle/), then follow links. Concepts follow the
[Open Knowledge Format](https://github.com/GoogleCloudPlatform/open-knowledge-format)."""

LINK = re.compile(r"\]\(([^)\s]+)\)")
FOOTNOTE_USE = re.compile(r"\[\^([^\]]+)\](?!:)")
PLACEHOLDER_ROW = re.compile(r"^\|\s*`(\[[A-Z0-9_/]+\])`", re.M)


def split_front(text: str) -> tuple[dict | None, str]:
    if not text.startswith("---\n"):
        return None, text
    try:
        _, front, body = text.split("---\n", 2)
    except ValueError:
        return None, text
    return yaml.safe_load(front) or {}, body


def concepts() -> list[Path]:
    return sorted(p for p in BRAIN.rglob("*.md") if p.name not in RESERVED)


def check_names(errors: list[str]) -> None:
    for p in list(BRAIN.rglob("*")) + list((ROOT / "templates").rglob("*")) + list((ROOT / "skills").rglob("*")):
        if " " in p.name or ("-" in p.name and p.suffix != ".py"):
            errors.append(f"{p.relative_to(ROOT)}: use _ in names, not spaces or hyphens")


def check_concepts(errors: list[str], warnings: list[str]) -> None:
    now = datetime.now(timezone.utc)
    for p in concepts():
        rel = p.relative_to(ROOT)
        meta, body = split_front(p.read_text())
        if meta is None:
            errors.append(f"{rel}: no YAML frontmatter")
            continue
        if not meta.get("type"):
            errors.append(f"{rel}: frontmatter has no `type`")
        for key in ("title", "description"):
            if not meta.get(key):
                warnings.append(f"{rel}: no `{key}`")
        source_ids = {s.get("id") for s in meta.get("sources") or [] if isinstance(s, dict)}
        for fid in set(FOOTNOTE_USE.findall(body)):
            if fid not in source_ids:
                errors.append(f"{rel}: footnote [^{fid}] has no matching sources[].id")
        stale = meta.get("stale_after")
        if stale:
            when = stale if isinstance(stale, datetime) else datetime.fromisoformat(str(stale).replace("Z", "+00:00"))
            if when.tzinfo is None:
                when = when.replace(tzinfo=timezone.utc)
            if now >= when:
                warnings.append(f"{rel}: stale since {stale}, re-verify it")
        if meta.get("type") == "Attested Computation":
            code = BRAIN / str(meta.get("computation", "")).lstrip("/")
            if not code.is_file():
                errors.append(f"{rel}: computation file {meta.get('computation')} not found")
            elif hashlib.sha256(code.read_bytes()).hexdigest() != meta.get("computation_sha256"):
                errors.append(f"{rel}: computation_sha256 doesn't match {meta['computation']}. Review the change, then update the hash")


def check_links(warnings: list[str]) -> None:
    for p in sorted(BRAIN.rglob("*.md")):
        for target in LINK.findall(p.read_text()):
            if re.match(r"^[a-z]+:", target) or target.startswith("#"):
                continue
            target = target.split("#")[0]
            dest = BRAIN / target.lstrip("/") if target.startswith("/") else (p.parent / target)
            if not dest.exists():
                warnings.append(f"{p.relative_to(ROOT)}: broken link {target}")


def check_indexes(errors: list[str]) -> None:
    for p in BRAIN.rglob("index.md"):
        meta, _ = split_front(p.read_text())
        if meta is not None and (p.parent != BRAIN or set(meta) - {"okf_version"}):
            errors.append(f"{p.relative_to(ROOT)}: index.md may only carry okf_version, and only at the bundle root")


def check_placeholders(errors: list[str]) -> None:
    sys.path.insert(0, str(Path(__file__).parent))
    from docx import Document
    from docx_tools import find_placeholders

    schema = BRAIN / "fields" / "project_fields.md"
    known = set(PLACEHOLDER_ROW.findall(schema.read_text()))
    for path in sorted((ROOT / "templates" / "word").glob("*.docx")):
        for ph in sorted(find_placeholders(Document(path)) - known):
            errors.append(f"{path.relative_to(ROOT)}: {ph} is not in brain/fields/project_fields.md")


def check_workbooks(errors: list[str]) -> None:
    """Every formula that reads the Setup sheet must point at a labelled row."""
    import openpyxl

    for path in sorted((ROOT / "templates" / "excel").glob("*.xlsx")):
        wb = openpyxl.load_workbook(path)
        if "Setup" not in wb.sheetnames:
            continue
        setup = wb["Setup"]
        for ws in wb:
            if ws.title == "Setup":
                continue
            for row in ws.iter_rows():
                for cell in row:
                    if not (isinstance(cell.value, str) and "Setup!" in cell.value):
                        continue
                    for ref in re.findall(r"Setup!\$?B\$?(\d+)", cell.value):
                        if not setup.cell(row=int(ref), column=1).value:
                            errors.append(f"{path.relative_to(ROOT)}: {ws.title}!{cell.coordinate} reads Setup!B{ref}, which has no label")


def write_indexes() -> None:
    dirs = sorted({p.parent for p in BRAIN.rglob("*.md")} | {BRAIN}, key=lambda d: len(d.parts), reverse=True)
    for d in dirs:
        lines = []
        subdirs = sorted(x for x in d.iterdir() if x.is_dir() and any(x.rglob("*.md")))
        docs = sorted(x for x in d.glob("*.md") if x.name not in RESERVED)
        if d == BRAIN:
            lines += ["---", f'okf_version: "{OKF_VERSION}"', "---", "", ROOT_INTRO, ""]
        if subdirs:
            lines += ["# Folders", ""]
            lines += [f"* [{TITLES.get(x.name, x.name.replace('_', ' ').capitalize())}]({x.name}/) - {FOLDERS.get(x.name, '')}".rstrip(" -") for x in subdirs]
            lines.append("")
        if docs:
            entries = []
            for x in docs:
                meta, _ = split_front(x.read_text())
                meta = meta or {}
                order = meta.get("phase_order") or meta.get("stage_order") or meta.get("order") or 0
                entries.append((order, meta.get("title", x.stem), x.name, meta.get("description", "")))
            lines += ["# Concepts", ""]
            lines += [f"* [{t}]({n}) - {desc}" for _, t, n, desc in sorted(entries)]
            lines.append("")
        (d / "index.md").write_text("\n".join(lines))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--write_indexes", action="store_true")
    args = parser.parse_args()
    if args.write_indexes:
        write_indexes()

    errors: list[str] = []
    warnings: list[str] = []
    check_names(errors)
    check_concepts(errors, warnings)
    check_indexes(errors)
    check_links(warnings)
    check_placeholders(errors)
    check_workbooks(errors)

    for w in warnings:
        print(f"warning: {w}")
    for e in errors:
        print(f"error: {e}")
    print(f"{len(concepts())} concepts, {len(errors)} errors, {len(warnings)} warnings")
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
